import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def create_synopsis():
    doc = docx.Document()
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = title.add_run("SYNOPSIS\n")
    run.bold = True
    run.font.size = Pt(16)
    
    # Header details
    doc.add_paragraph("Title of the Project: CropCare AI: Detection of Crop Disease Using CNN and RAG-Enabled Decision Support")
    doc.add_paragraph("USN: _______________________")
    doc.add_paragraph("Name: ______________________")
    doc.add_paragraph("Guide Name: _________________")
    
    doc.add_paragraph()
    
    # Introduction
    heading = doc.add_heading('Introduction', level=1)
    
    intro_text1 = (
        "Agriculture is the backbone of the economy, yet it constantly faces severe threats from various plant diseases "
        "caused by pathogens and adverse environmental factors. These diseases, such as leaf spot, blight, rust, powdery mildew, "
        "bacterial infections, and viral diseases (e.g., mosaic and leaf-curl symptoms), can drastically reduce crop health, yield, and "
        "ultimately the farmers' livelihood if not diagnosed and treated in a timely manner. Traditional methods of disease "
        "identification rely heavily on the manual inspection by human experts, which is often time-consuming, expensive, and prone to error.\n"
    )
    intro_text2 = (
        "To address this critical issue, this project introduces 'CropCare AI', an explainable crop disease detection and decision-support "
        "system designed to assist farmers and agricultural workers in identifying diseases directly from crop leaf images. "
        "The system leverages state-of-the-art Deep Learning technologies, specifically a Convolutional Neural Network (CNN) based on the VGG-16 architecture, "
        "coupled with a custom Spatial Attention Mechanism. This approach allows the model to intelligently focus on the most relevant diseased regions (Region of Interest) "
        "of a leaf while ignoring background noise, thereby significantly improving the accuracy of disease classification.\n"
    )
    intro_text3 = (
        "Beyond simple classification, CropCare AI integrates a Retrieval-Augmented Generation (RAG) framework to provide evidence-based, grounded decision support. "
        "When a disease is detected, the system retrieves relevant information from a curated agricultural knowledge base. It generates a comprehensive, downloadable PDF report "
        "detailing the disease symptoms, agricultural precautions, preventive measures, and management solutions. Furthermore, the platform bridges the gap between farmers "
        "and professionals by offering a built-in consultation module where users can locate agricultural experts, view their expertise, and book appointment time slots. "
        "By combining advanced computer vision, generative AI, and a seamless user experience, CropCare AI provides a holistic, practical workflow from early diagnosis to expert guidance."
    )
    
    p = doc.add_paragraph(intro_text1)
    p = doc.add_paragraph(intro_text2)
    p = doc.add_paragraph(intro_text3)
    
    # Problem Statement / Objectives
    doc.add_heading('Problem Statement / Objectives', level=1)
    doc.add_paragraph("Problem Statement:", style='List Bullet')
    doc.add_paragraph("Crop diseases lead to significant agricultural losses, but farmers often lack immediate access to accurate diagnostic tools and expert advice, leading to delayed treatments or the misuse of chemical pesticides.", style='List Bullet 2')
    doc.add_paragraph("Objectives:", style='List Bullet')
    doc.add_paragraph("To develop an automated, highly accurate CNN-based image processing pipeline to detect crop diseases from leaf images.", style='List Bullet 2')
    doc.add_paragraph("To implement an attention mechanism and Region of Interest (ROI) extraction to improve classification reliability and handle uncertain predictions safely via confidence thresholds.", style='List Bullet 2')
    doc.add_paragraph("To integrate a RAG-based AI system that provides actionable, context-aware agricultural advice based on the detected disease.", style='List Bullet 2')
    doc.add_paragraph("To create a responsive web platform that generates downloadable PDF diagnostic reports.", style='List Bullet 2')
    doc.add_paragraph("To establish a consultation network allowing farmers to discover and book appointments with agricultural advisors.", style='List Bullet 2')
    
    # Platforms and Tools used
    doc.add_heading('Platforms and Tools used', level=1)
    doc.add_paragraph("Platform: Web-based Application (Responsive Web App)", style='List Bullet')
    doc.add_paragraph("Tools:", style='List Bullet')
    doc.add_paragraph("Front-End: HTML5, CSS3, Vanilla JavaScript", style='List Bullet 2')
    doc.add_paragraph("Back-End: Python (Flask Framework)", style='List Bullet 2')
    doc.add_paragraph("Machine Learning: PyTorch, TorchVision (VGG-16 with Attention Mechanism)", style='List Bullet 2')
    doc.add_paragraph("Image Processing: OpenCV (ROI Extraction)", style='List Bullet 2')
    doc.add_paragraph("Generative AI & Decision Support: Google Gemini API (RAG-based Retrieval)", style='List Bullet 2')
    doc.add_paragraph("Database: MongoDB", style='List Bullet 2')
    doc.add_paragraph("Data Set: PlantVillage Dataset (and curated agricultural knowledge bases for RAG)", style='List Bullet')
    doc.add_paragraph("Algorithm Used: Convolutional Neural Networks (CNN) with Spatial Attention, Retrieval-Augmented Generation (RAG)", style='List Bullet')

    # Methodology
    doc.add_heading('Methodology', level=1)
    p_meth = doc.add_paragraph("The methodology follows a structured approach to achieve the project objectives:")
    
    doc.add_heading('a. Application Based Project', level=2)
    doc.add_paragraph("1. Data Collection & Preprocessing: Gather leaf images from the PlantVillage dataset. Perform image preprocessing, including resizing, normalization, and OpenCV-based ROI extraction to isolate the green leaf from backgrounds.", style='List Number')
    doc.add_paragraph("2. Model Development: Train a VGG-16 CNN integrated with a custom attention block using PyTorch to classify the diseases and output a confidence score.", style='List Number')
    doc.add_paragraph("3. AI Integration: Implement the Google Gemini API to act as a RAG layer, dynamically fetching curated disease management strategies based on the CNN's prediction.", style='List Number')
    doc.add_paragraph("4. Web Application Development: Build a Flask-based backend with MongoDB for user and advisor data storage. Develop a responsive Glassmorphism UI for seamless interaction.", style='List Number')
    doc.add_paragraph("5. Feature Implementation: Add role-based dashboards, automated PDF report generation (using ReportLab), and an interactive advisor appointment booking system with double-booking prevention and map integration.", style='List Number')

    doc.add_heading('Hardware Requirements (Development & Deployment):', level=3)
    doc.add_paragraph("Processor: Intel Core i5 or equivalent (Minimum); Intel Core i7 / GPU recommended for deep learning training.", style='List Bullet')
    doc.add_paragraph("RAM: 8 GB Minimum (16 GB Recommended).", style='List Bullet')
    doc.add_paragraph("Storage: 256 GB SSD (for fast data processing of image datasets).", style='List Bullet')
    doc.add_paragraph("Network: Active Internet connection for API integrations and MongoDB Atlas.", style='List Bullet')

    doc.add_heading('Software Requirements (Development & Deployment):', level=3)
    doc.add_paragraph("Operating System: Windows 10/11, macOS, or Linux.", style='List Bullet')
    doc.add_paragraph("Programming Languages: Python 3.9+, HTML5, CSS3, JavaScript (ES6+).", style='List Bullet')
    doc.add_paragraph("Frameworks & Libraries: Flask, PyTorch, TorchVision, OpenCV, ReportLab, Leaflet.js.", style='List Bullet')
    doc.add_paragraph("Database: MongoDB (Local or Atlas).", style='List Bullet')
    doc.add_paragraph("IDE/Editor: VS Code, PyCharm, or Antigravity IDE.", style='List Bullet')

    doc.add_heading('b. Benefits of the project for the society', level=2)
    doc.add_paragraph("1. Empowers Farmers: Provides immediate, accessible, and accurate disease diagnostics directly to farmers' smartphones or computers, reducing reliance on scarce human experts.", style='List Bullet')
    doc.add_paragraph("2. Enhances Food Security: Early detection and proper management of crop diseases prevent massive yield losses, ensuring a more stable food supply.", style='List Bullet')
    doc.add_paragraph("3. Promotes Sustainable Agriculture: By providing precise, evidence-based management solutions, the system reduces the indiscriminate use of harmful chemical pesticides, protecting the environment.", style='List Bullet')
    doc.add_paragraph("4. Bridges the Expert Gap: Facilitates direct communication and appointment booking between rural farmers and certified agricultural advisors, democratizing access to expert knowledge.", style='List Bullet')

    doc.save("Project_Synopsis_V2.docx")
    print("Synopsis generated successfully.")

if __name__ == "__main__":
    create_synopsis()
